import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def create_docx(md_file, docx_file):
    doc = docx.Document()
    
    # Thiết lập font mặc định chuẩn đại học (Times New Roman, 13pt)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)

    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_text = ""

    for line in lines:
        line = line.rstrip('\n')
        
        # Bỏ qua dòng HTML anchors (<a id=...>) và horizontal rules (---)
        if line.startswith('<a id') or line == '---':
            continue
            
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_text = ""
            else:
                in_code_block = False
                p = doc.add_paragraph()
                run = p.add_run(code_text.strip())
                run.font.name = 'Courier New'
                run.font.size = Pt(10)
            continue
            
        if in_code_block:
            code_text += line + "\n"
            continue
            
        # Xử lý Headings và bỏ qua markdown links dạng [Text](#link)
        clean_line = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', line)
        clean_line = re.sub(r'\*\*(.*?)\*\*', r'\1', clean_line) # Bỏ markdown bold
        clean_line = re.sub(r'\*(.*?)\*', r'\1', clean_line) # Bỏ markdown italic
        
        if clean_line.startswith('# '):
            doc.add_heading(clean_line[2:], level=1)
        elif clean_line.startswith('## '):
            doc.add_heading(clean_line[3:], level=2)
        elif clean_line.startswith('### '):
            doc.add_heading(clean_line[4:], level=3)
        elif clean_line.strip().startswith('- '):
            doc.add_paragraph(clean_line.strip()[2:], style='List Bullet')
        elif clean_line.strip() == '':
            pass
        else:
            doc.add_paragraph(clean_line)

    doc.save(docx_file)
    print(f"Đã tạo thành công {docx_file}")

if __name__ == '__main__':
    create_docx('TAI_LIEU_THIET_KE_V2.md', 'TAI_LIEU_THIET_KE_V2.docx')
